"""Unit tests for document parsers."""

import pytest
import tempfile
import os
import pymupdf as fitz
from openpyxl import Workbook

from app.ingestion.parsers.markdown_parser import MarkdownParser
from app.ingestion.parsers.txt_parser import TXTParser
from app.ingestion.parsers.csv_parser import CSVParser
from app.ingestion.parsers.pdf_parser import PDFParser
from app.ingestion.parsers.xlsx_parser import XLSXParser


@pytest.mark.asyncio
async def test_pdf_parser():
    parser = PDFParser()
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "ASK Financial Holdings Portfolio Overview 2026.")
    
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_path = f.name
    doc.save(temp_path)
    doc.close()

    try:
        parsed = await parser.parse(temp_path, "test.pdf")
        assert len(parsed.sections) >= 1
        assert "ASK Financial" in parsed.sections[0].content
        assert parsed.metadata["pages"] == 1
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_xlsx_and_xlsm_parser():
    parser = XLSXParser()
    assert ".xlsm" in parser.supported_extensions()
    assert ".xlsx" in parser.supported_extensions()

    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Quarter", "Revenue", "Margin"])
    ws.append(["Q1", "500000", "0.25"])
    ws.append(["Q2", "620000", "0.28"])

    with tempfile.NamedTemporaryFile(suffix=".xlsm", delete=False) as f:
        temp_path = f.name
    wb.save(temp_path)
    wb.close()

    try:
        parsed = await parser.parse(temp_path, "financials.xlsm")

        # One section per data row, not one blob per sheet: a sheet flattened
        # into a single block gets split on character count, and every chunk
        # after the first loses the header row — leaving bare values with
        # nothing saying which column each belongs to.
        assert len(parsed.sections) == 2

        first = parsed.sections[0]
        assert first.section == "Revenue \u00b7 Q1"
        # Each row carries its own headers, which is what lets a question about
        # one row match that row instead of a slab of the sheet.
        assert "Quarter: Q1" in first.content
        assert "Revenue: 500000" in first.content
        assert first.metadata["sheet"] == "Revenue"
        assert first.metadata["row_number"] == 2

        assert parsed.sections[1].section == "Revenue \u00b7 Q2"
        assert "Revenue: 620000" in parsed.sections[1].content

        assert len(parsed.tables) == 1
        assert parsed.tables[0]["headers"] == ["Quarter", "Revenue", "Margin"]
        assert len(parsed.tables[0]["rows"]) == 2
        assert parsed.metadata["format"] == "xlsm"
        assert parsed.metadata["sheets"] == 1
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_markdown_parser():
    parser = MarkdownParser()
    content = """# Company Handbook

## Leave Policy
Employees get 20 days of paid annual leave.

## Health Insurance
Comprehensive medical coverage is provided.
"""

    with tempfile.NamedTemporaryFile(suffix=".md", delete=False, mode="w", encoding="utf-8") as f:
        f.write(content)
        temp_path = f.name

    try:
        parsed = await parser.parse(temp_path, "handbook.md")
        assert len(parsed.sections) >= 2
        sections = [s.section for s in parsed.sections if s.section]
        assert "Leave Policy" in sections
        assert "Health Insurance" in sections
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_txt_parser():
    parser = TXTParser()
    content = """First paragraph of policy document.

Second paragraph with more details."""

    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
        f.write(content)
        temp_path = f.name

    try:
        parsed = await parser.parse(temp_path, "policy.txt")
        assert len(parsed.sections) == 2
        assert "First paragraph" in parsed.sections[0].content
        assert "Second paragraph" in parsed.sections[1].content
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_csv_parser():
    parser = CSVParser()
    content = """Department,Headcount,Budget
Engineering,45,1200000
Finance,12,450000
Human Resources,8,250000
"""

    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w", encoding="utf-8") as f:
        f.write(content)
        temp_path = f.name

    try:
        parsed = await parser.parse(temp_path, "headcount.csv")
        assert len(parsed.sections) >= 1
        assert "Engineering" in parsed.sections[0].content
        assert len(parsed.tables) == 1
        assert parsed.tables[0]["headers"] == ["Department", "Headcount", "Budget"]
        assert len(parsed.tables[0]["rows"]) == 3
    finally:
        os.unlink(temp_path)
